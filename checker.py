#!/usr/bin/python3

from email.message import EmailMessage
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.wait import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from dotenv import load_dotenv
import os
import time
import re
import datetime
import traceback
import smtplib
import json

# A personal environment variable file to get the email's app password
load_dotenv("/home/mimis/.mail.env")

today = datetime.datetime.today()
print(today)

highs = []
prices = []
excel_path = "/4TB_WD/Mimis/Prices/prices.xlsx"
price_pattern = r"[0-9]+\.[0-9]+$" # This matches what a price should look like, so that we never get something inconsistent if the site changes or something

# Getting the links from the excel file
links = []
excel = load_workbook(excel_path)
sheet = excel.active # Opening the excel sheet and NOT closing it until the bottom of the script
row = 2

while sheet[f"C{row}"].value != None:
    links.append(sheet[f"C{row}"].value)
    highs.append(sheet[f"E{row}"].value)
    row += 1

def send_mail(to, fund_code, link_index, maxPercent):
    roundedPercent = round(maxPercent, 1) # Rounded to the first decimal place

    mail = EmailMessage()
    mail["From"] = "dimitris_the_programmer@yahoo.com"
    mail["To"] = to
    mail["Subject"] = f"Fund { fund_code } fell another 5% tier (or more)"
    mail.set_content(f"It's now sitting at {roundedPercent}% of it's maximum. \nCheck it out at { links[link_index] }")

    server = smtplib.SMTP_SSL("smtp.mail.yahoo.com", 465)
    server.login(mail["From"], os.getenv("MAIL_PASS")) # The email's specific app password from the environment variable file loaded above
    server.send_message(mail)
    server.quit()

# Ignoring superficial logs, running the browser with no graphics (but NOT headless, unless you like captchas), and other minor details you definitely know (!).
options = webdriver.ChromeOptions()
options.add_argument('--log-level=3')
options.add_argument("--disable-extensions")
options.add_argument("--disable-gpu")
options.add_argument("--disable-infobars")
options.add_argument("--disable-notifications")
options.add_argument("--no-sandbox")
options.add_argument("--disable-dev-shm-usage")

service = Service("/usr/bin/chromedriver")

driver = webdriver.Chrome(service=service, options=options)
wait = WebDriverWait(driver, 40)

for link in links:
    try:
        driver.get(link)
        time.sleep(10)

        if "blackrock" in link:
            # Accept necessary cookies
            locator = (By.ID, "onetrust-reject-all-handler")
            necessary = wait.until(EC.presence_of_element_located(locator))
            necessary.click()

            # Accept terms & conditions
            locator = (By.ID, "terms-and-conditions-masthead-accept")
            accept = wait.until(EC.presence_of_element_located(locator))
            accept.click()

            time.sleep(2)

            # First search for the fund because the whole link redirects to the product list, not the specific one (site's problem)
            locator = (By.ID, "screener-keyword-search-input")
            inputfield = wait.until(EC.presence_of_element_located(locator))
            inputfield.send_keys("LU0171289902") # The fund's code

            time.sleep(2)

            # Now get the nav's element from within a div within a td element with these classes
            locator = (By.CSS_SELECTOR, "td.mat-mdc-cell.mdc-data-table__cell.cdk-cell.cdk-column-navAmount.mat-column-navAmount.numeric-cell div")
            elem = wait.until(EC.presence_of_element_located(locator))

            nav = elem.text

            if re.search(price_pattern, nav):
                prices.append(nav)
            else:
                prices.append(-1)

        elif "jpmorgan" in link:
            # Only after those buttons are pressed does the site load the NAV.
            # Also, you are only asked to press it once (doesn't appear in the second visit). ie000u9j8hx9 is the code of the second fund from jpmorgan.
            if not "ie000u9j8hx9" in link:                                                                                     ####### IF IT STOPS AGAIN BECAUSE OF INTERCEPTION, CONSIDER REMOVING THIS
                try: # For some reason it may not ask you to stay or not on the UK site
                    locator = (By.CSS_SELECTOR, "button.cta-btn.secondary")
                    uk_button = wait.until(EC.presence_of_element_located(locator))
                    uk_button.click()
                except Exception:
                    print("JP Morgan UK button not found bobo")

                # As of writing this, it looks like they (temporarily?) added a feedback button which pops up above the cookie buttons and intercepts the click below (causing an error, which showed what intercepted the click)
                # So I'll click "No" and put it in a "try" block in case they remove this.
                try:
                    locator = (By.ID, "idbtnCancel")
                    NoFeedbackButton = wait.until(EC.presence_of_element_located(locator))
                    NoFeedbackButton.click()
                except Exception:
                    print("Feedback button in JPMorgan not found. Test it once manually, and it must be safe to remove the code above this print.")

                # Check for the UK button again (in case they were given in a different order. Avoids click interceptions)
                try: # For some reason it may not ask you to stay or not on the UK site
                    locator = (By.CSS_SELECTOR, "button.cta-btn.secondary")
                    uk_button = wait.until(EC.presence_of_element_located(locator))
                    uk_button.click()
                except Exception:
                    print("JP Morgan UK button not found bobo (twice)")


                locator = (By.ID, "onetrust-reject-all-handler")
                reject_btn = wait.until(EC.presence_of_element_located(locator))
                reject_btn.click()

                locator = (By.CSS_SELECTOR, "a.accept")
                accept_terms = wait.until(EC.presence_of_element_located(locator))
                accept_terms.click()

            # Then just get the NAV by the element with the argument data-testid set to [this]
            locator = (By.CSS_SELECTOR, '[data-testid="quick-stat-nav-value"]')
            nav_elem = wait.until(EC.presence_of_element_located(locator))

            nav = nav_elem.text[4:] # Also removing the "EUR " from the text, to be left with a numerical price

            if re.search(price_pattern, nav):
                prices.append(nav)
            else:
                prices.append(-1)

        elif "markets.ft" in link:
            # FINISHING THIS ONE, IS NOT AVAILABLE HERE (GETTING LIQUIDATED).
            print("LU1775978304 is no longer available (was on markets.ft). Skipping...")
            prices.append(-1)
            continue

            # Like in BlackRock, there are many elements with the same tag and class as the one with the NAV price. Working similarly.
            locator = (By.CSS_SELECTOR, "span.mod-ui-data-list__value")
            elems = wait.until(EC.presence_of_all_elements_located(locator))

            nav = elems[0].text # NAV is the 1st span with this class

            if re.search(price_pattern, nav):
                prices.append(nav)
            else:
                prices.append(-1)

        elif "franklintempleton" in link:
            # No questions asked, just loading the site & getting the price
            locator = (By.CSS_SELECTOR, "p.summary-item__value.summary-item__value--basic")
            nav_elem = wait.until(EC.presence_of_element_located(locator))

            nav = nav_elem.text[1:] # Excluding "€" from the NAV text

            if re.search(price_pattern, nav):
                prices.append(nav)
            else:
                prices.append(-1)

        elif "tradingview" in link:
            # Same, just get the price
            locator = (By.CSS_SELECTOR, ".last-KLji300y.js-symbol-last")
            price_elem = wait.until(EC.presence_of_element_located(locator))

            price = price_elem.text
            if re.search(price_pattern, price):
                prices.append(price)
            else:
                prices.append(-1)

    except Exception:
        traceback.print_exc()
        prices.append(-1)

def getTier(maxPercent):
    tier = int( (maxPercent // 10) * 10) # The tens digit of maxPercent (without int(), the result is an ugly float, with the decimal always being 0)

    if maxPercent - tier < 5:
        return tier

    return tier + 5

# Finally writing the prices & messing with the colors
normal_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid") # White background cell fill
missing_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid") # Yellow background cell fill

with open("/home/mimis/Projects/FundPriceChecker/prevMaxPercentTiers.json") as f:
    prevMaxPercentTier = json.load(f)

for i, price in enumerate(prices, start=2):
    if price != -1:
        sheet[f"A{i}"].fill = normal_fill
        sheet[f"A{i}"] = float(price)
        sheet[f"D{i}"] = '=VALUE("' + today.strftime("%d/%m/%Y") + '")' # Too many quotes to use an f-string...

        # Increase the max if the price exceeds it. Column E holds the max prices.
        max = sheet[f"E{i}"].value
        code = sheet[f"B{i}"].value

        if (max == None or max == "") or float(price) >= max: # Also setting it up like this if it's new and it doesnt have a max
            sheet[f"E{i}"] = float(price)
            prevMaxPercentTier[code] = 95

        else: # Also handles the case where there is no prevMaxPercentTier. No matter what, it will be set, either according to max (even if it's not set), or the current maxPercent.
            maxPercent = (float(price) / max) * 100

            if maxPercent < prevMaxPercentTier.get(code, 0): # If the price as a percentage of max drops below the 5% tiers contained in prevMaxPercentTier, drop the new tier and send the mail
                send_mail("dimitris_the_programmer@yahoo.com", code, i-2, maxPercent)
                send_mail("a123xxsp@gmail.com", code, i-2, maxPercent)

            prevMaxPercentTier[code] = getTier(maxPercent)

    else:
        sheet[f"A{i}"].fill = missing_fill # We don't change the text value, just inform that it is a past price & something went wrong here or on the site.

excel.save(excel_path)

with open("/home/mimis/Projects/FundPriceChecker/prevMaxPercentTiers.json", "w") as f:
    json.dump(prevMaxPercentTier, f, indent=4)

print("")
